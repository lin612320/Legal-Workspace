#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Markdown → DOCX 转换（法元文档用）

用途：仓库里 docs/*.md 与 *.docx 成对存在，本脚本用于把中文 Markdown 同步为 Word，
避免"改了 md、docx 还是旧的"这类文档漂移。

用法：
    py -3 scripts/md2docx.py docs/项目计划书.md docs/项目计划书.docx
    py -3 scripts/md2docx.py docs/Project-Plan-EN.md docs/Project-Plan-EN.docx

支持：标题（#~####）、段落、无序/有序列表、表格（含表头分隔行）、引用块（>）、
分隔线（---）、行内 **加粗** / `代码` / [链接](url) / <自动链接>、中文正文字体设置。

不追求完整 CommonMark：只覆盖本项目文档实际使用的语法子集，未知语法按段落原文输出。
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

CJK_FONT = "Microsoft YaHei"
LATIN_FONT = "Segoe UI"
MONO_FONT = "Consolas"
QUOTE_COLOR = RGBColor(0x55, 0x5F, 0x70)
CODE_COLOR = RGBColor(0x8A, 0x1F, 0x18)

INLINE_RE = re.compile(
    r"(\*\*.+?\*\*)"      # 加粗
    r"|(`[^`]+`)"          # 行内代码
    r"|(\[[^\]]+\]\([^)]+\))"  # 链接
    r"|(<https?://[^>]+>)"  # 自动链接
)


def set_run_font(run, name: str) -> None:
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), CJK_FONT)


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.size = Pt(10.5)
    normal.font.name = LATIN_FONT
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.35
    for style_name in ("Heading 1", "Heading 2", "Heading 3", "Heading 4"):
        try:
            st = doc.styles[style_name]
        except KeyError:
            continue
        st.font.name = LATIN_FONT
        if st.element.rPr is not None:
            rfonts = st.element.rPr.find(qn("w:rFonts"))
            if rfonts is None:
                rfonts = OxmlElement("w:rFonts")
                st.element.rPr.append(rfonts)
            rfonts.set(qn("w:eastAsia"), CJK_FONT)


def add_inline(paragraph, text: str, *, bold_all: bool = False, color=None, italic: bool = False) -> None:
    """按行内语法切分并写入 runs。"""
    pos = 0
    for m in INLINE_RE.finditer(text):
        if m.start() > pos:
            run = paragraph.add_run(text[pos:m.start()])
            run.bold = bold_all
            run.italic = italic
            if color is not None:
                run.font.color.rgb = color
            set_run_font(run, LATIN_FONT)
        token = m.group(0)
        if token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
            run.italic = italic
            if color is not None:
                run.font.color.rgb = color
            set_run_font(run, LATIN_FONT)
        elif token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.color.rgb = CODE_COLOR
            set_run_font(run, MONO_FONT)
        elif token.startswith("["):
            label, url = re.match(r"\[([^\]]+)\]\(([^)]+)\)", token).groups()
            run = paragraph.add_run(label)
            run.bold = bold_all
            run.font.color.rgb = RGBColor(0x1A, 0x4F, 0x9C)
            set_run_font(run, LATIN_FONT)
            if url and url != label:
                tail = paragraph.add_run(f"（{url}）")
                tail.font.size = Pt(8.5)
                tail.font.color.rgb = QUOTE_COLOR
                set_run_font(tail, LATIN_FONT)
        else:  # <https://...>
            run = paragraph.add_run(token[1:-1])
            run.font.color.rgb = RGBColor(0x1A, 0x4F, 0x9C)
            set_run_font(run, LATIN_FONT)
        pos = m.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        run.bold = bold_all
        run.italic = italic
        if color is not None:
            run.font.color.rgb = color
        set_run_font(run, LATIN_FONT)


def add_hr(doc: Document) -> None:
    p = doc.add_paragraph()
    ppr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:color"), "BBBBBB")
    borders.append(bottom)
    ppr.append(borders)


def split_row(line: str) -> list[str]:
    cells = line.strip().strip("|").split("|")
    return [c.strip() for c in cells]


def is_sep_row(line: str) -> bool:
    return bool(re.match(r"^\|[\s:\-|]+\|$", line.strip()))


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=0, cols=cols)
    table.style = "Table Grid"
    for i, cells in enumerate(rows):
        row = table.add_row()
        for j in range(cols):
            cell_text = cells[j] if j < len(cells) else ""
            para = row.cells[j].paragraphs[0]
            para.paragraph_format.space_after = Pt(2)
            add_inline(para, cell_text, bold_all=(i == 0))
            for run in para.runs:
                run.font.size = Pt(9)


def convert(md_path: Path, docx_path: Path) -> tuple[int, int]:
    lines = md_path.read_text(encoding="utf-8").replace("\r\n", "\n").split("\n")
    doc = Document()
    configure_styles(doc)

    i = 0
    n_tables = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # 表格
        if stripped.startswith("|") and i + 1 < len(lines) and is_sep_row(lines[i + 1]):
            rows = [split_row(stripped)]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(split_row(lines[i]))
                i += 1
            add_table(doc, rows)
            n_tables += 1
            doc.add_paragraph()
            continue

        # 标题
        m = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            doc.add_heading(m.group(2).strip(), level=min(level, 4))
            i += 1
            continue

        # 分隔线
        if re.match(r"^-{3,}$", stripped):
            add_hr(doc)
            i += 1
            continue

        # 引用块（连续行合并）
        if stripped.startswith(">"):
            block = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                block.append(lines[i].strip().lstrip(">").strip())
                i += 1
            text = " ".join(t for t in block if t)
            if not text:
                continue
            para = doc.add_paragraph()
            para.paragraph_format.left_indent = Inches(0.25)
            para.paragraph_format.space_before = Pt(2)
            add_inline(para, text, italic=True, color=QUOTE_COLOR)
            for run in para.runs:
                run.font.size = Pt(9.5)
            continue

        # 列表
        m = re.match(r"^([-*+])\s+(.*)$", stripped)
        if m:
            para = doc.add_paragraph(style="List Bullet")
            add_inline(para, m.group(2))
            i += 1
            continue
        m = re.match(r"^(\d+)[.)]\s+(.*)$", stripped)
        if m:
            para = doc.add_paragraph(style="List Number")
            add_inline(para, m.group(2))
            i += 1
            continue

        # 空行
        if not stripped:
            i += 1
            continue

        # 普通段落
        para = doc.add_paragraph()
        add_inline(para, stripped)
        i += 1

    doc.save(str(docx_path))
    return len(lines), n_tables


def main() -> int:
    if len(sys.argv) != 3:
        print(__doc__)
        return 2
    md_path = Path(sys.argv[1])
    docx_path = Path(sys.argv[2])
    if not md_path.exists():
        print(f"源文件不存在：{md_path}")
        return 1
    lines, tables = convert(md_path, docx_path)
    size_kb = docx_path.stat().st_size / 1024
    print(f"已生成 {docx_path}（{lines} 行 Markdown → {tables} 张表格，{size_kb:.1f} KB）")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
